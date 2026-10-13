"""Compile the Beamer deck and create a pixel-faithful PowerPoint from it."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PRESENTATION_DIR = ROOT / "presentation"
TEX_PATH = PRESENTATION_DIR / "analyst_recommendation_anomaly.tex"
PDF_PATH = PRESENTATION_DIR / "analyst_recommendation_anomaly.pdf"
PPTX_PATH = PRESENTATION_DIR / "analyst_recommendation_anomaly.pptx"
RENDER_DIR = PRESENTATION_DIR / "rendered"


def require(program: str) -> str:
    resolved = shutil.which(program)
    if resolved is None:
        raise RuntimeError(f"Required program is not installed: {program}")
    return resolved


def build() -> tuple[Path, Path]:
    tectonic = require("tectonic")
    pdftoppm = require("pdftoppm")
    PRESENTATION_DIR.mkdir(parents=True, exist_ok=True)
    RENDER_DIR.mkdir(parents=True, exist_ok=True)

    subprocess.run(
        [tectonic, "--outdir", str(PRESENTATION_DIR), str(TEX_PATH)],
        cwd=ROOT,
        check=True,
    )
    for image in RENDER_DIR.glob("slide-*.png"):
        image.unlink()
    subprocess.run(
        [
            pdftoppm,
            "-png",
            "-r",
            "180",
            str(PDF_PATH),
            str(RENDER_DIR / "slide"),
        ],
        check=True,
    )

    images = sorted(RENDER_DIR.glob("slide-*.png"))
    if not images:
        raise RuntimeError("No Beamer pages were rendered")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for image in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
    prs.save(PPTX_PATH)
    return PDF_PATH, PPTX_PATH


if __name__ == "__main__":
    pdf, pptx = build()
    print(pdf)
    print(pptx)
